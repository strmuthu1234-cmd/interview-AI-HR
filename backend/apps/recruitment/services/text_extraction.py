"""Extract plain text from PDF (PyMuPDF) and DOCX (python-docx) resumes."""
import io
import logging

from django.conf import settings

log = logging.getLogger(__name__)

MIN_TEXT_CHARS = 80  # below this a PDF is treated as scanned / image-only


class ExtractionError(Exception):
    pass


def extract_pdf(data: bytes) -> str:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF: {exc}") from exc
    if doc.needs_pass:
        raise ExtractionError("PDF is password protected.")
    if doc.page_count > 15:
        raise ExtractionError("PDF has too many pages for a resume (limit 15).")
    parts = [page.get_text("text") for page in doc]
    text = "\n".join(parts).strip()
    if len(text) < MIN_TEXT_CHARS and settings.OCR_ENABLED:
        text = _ocr_pdf(doc) or text
    return text


def _ocr_pdf(doc) -> str:
    """Optional OCR fallback (requires Tesseract; enabled with OCR_ENABLED=true)."""
    try:
        out = []
        for page in doc:
            tp = page.get_textpage_ocr(full=True)
            out.append(page.get_text("text", textpage=tp))
        return "\n".join(out).strip()
    except Exception as exc:  # Tesseract missing etc.
        log.warning("OCR fallback failed: %s", exc)
        return ""


def extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in cells:
                    cells.append(t)
            if cells:
                lines.append("  |  ".join(cells))
    return "\n".join(lines).strip()


def extract_text(file_type: str, data: bytes) -> str:
    text = extract_pdf(data) if file_type == "pdf" else extract_docx(data)
    text = text.replace("\x00", "")
    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "Little or no text could be extracted (the resume may be a scanned image). "
            "Enable OCR (OCR_ENABLED=true with Tesseract installed) or ask for a text-based resume."
        )
    return text[:60000]
